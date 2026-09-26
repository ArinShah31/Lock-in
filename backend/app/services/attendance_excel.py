"""Persistent per-classroom attendance workbook (DB stays the source of truth)."""

import logging
import os
import re
import threading
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.services.attendance import Snapshot, load_snapshot

logger = logging.getLogger(__name__)

ATTENDANCE_DIR = Path("uploads/attendance")
ATTENDANCE_DIR.mkdir(parents=True, exist_ok=True)

_locks: dict[int, threading.Lock] = {}
_locks_guard = threading.Lock()

_HEADER_FILL = PatternFill("solid", fgColor="E8EDF5")
_BOLD = Font(bold=True)


def workbook_path(classroom_id: int) -> Path:
    return ATTENDANCE_DIR / f"classroom_{classroom_id}.xlsx"


def _lock_for(classroom_id: int) -> threading.Lock:
    with _locks_guard:
        return _locks.setdefault(classroom_id, threading.Lock())


def _sheet_title(name: str, used: set[str]) -> str:
    base = re.sub(r"[\[\]\*\?/\:]", "-", name).strip() or "Subject"
    base = base[:28]
    title, n = base, 2
    while title.lower() in used or title.lower() == "summary":
        title = f"{base}-{n}"
        n += 1
    used.add(title.lower())
    return title


def _fmt_pct(value: float | None):
    return "" if value is None else value


def build_workbook(snap: Snapshot) -> Workbook:
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    used: set[str] = {"summary"}

    for subject in snap.subjects:
        ws = wb.create_sheet(_sheet_title(subject.code or subject.name, used))
        sessions = [s for s in snap.sessions if s.subject_id == subject.id]
        header = ["Student", "Email"]
        for s in sessions:
            label = s.held_on.isoformat()
            header.append(f"{label} {s.title}".strip())
        header += ["Attended", "Held", "%"]
        ws.append(header)
        for cell in ws[1]:
            cell.font = _BOLD
            cell.fill = _HEADER_FILL
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
        for student in snap.students:
            row = [student.full_name, student.email]
            attended = held = 0
            for s in sessions:
                mark = snap.marks.get((s.id, student.id))
                if mark is None:
                    row.append("")
                    continue
                held += 1
                attended += 1 if mark else 0
                row.append("P" if mark else "A")
            pct = round(attended * 100.0 / held, 2) if held else None
            row += [attended, held, _fmt_pct(pct)]
            ws.append(row)
        ws.column_dimensions["A"].width = 26
        ws.column_dimensions["B"].width = 30
        for i in range(3, len(header) + 1):
            ws.column_dimensions[get_column_letter(i)].width = 14
        ws.freeze_panes = "C2"

    header = ["Student", "Email", *[s.name for s in snap.subjects], "Final %"]
    summary.append(header)
    for cell in summary[1]:
        cell.font = _BOLD
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    for student in snap.students:
        per_subject, final = snap.student_stats(student.id)
        summary.append(
            [
                student.full_name,
                student.email,
                *[_fmt_pct(per_subject[s.id]["percent"]) for s in snap.subjects],
                _fmt_pct(final),
            ]
        )
    summary.column_dimensions["A"].width = 26
    summary.column_dimensions["B"].width = 30
    for i in range(3, len(header) + 1):
        summary.column_dimensions[get_column_letter(i)].width = 16
    summary.freeze_panes = "C2"
    return wb


def rebuild_workbook(db: Session, classroom_id: int) -> Path | None:
    """Rewrite the classroom's persistent .xlsx. Never raises: logs on failure."""
    path = workbook_path(classroom_id)
    tmp = path.with_suffix(".xlsx.tmp")
    try:
        with _lock_for(classroom_id):
            wb = build_workbook(load_snapshot(db, classroom_id))
            wb.save(tmp)
            os.replace(tmp, path)
        return path
    except Exception:  # noqa: BLE001
        logger.exception("Could not rebuild attendance workbook for classroom %s", classroom_id)
        try:
            tmp.unlink(missing_ok=True)
        except OSError:
            pass
        return None
