import tempfile
import unittest
from pathlib import Path
from unittest import mock

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.core.database import Base, get_db
from app.core.security import create_access_token, get_password_hash
from app.main import app
from app.models.attendance import AttendanceRecord, AttendanceSession
from app.models.classroom import (
    Classroom,
    ClassroomStudent,
    ClassroomTeacher,
    MembershipStatus,
)
from app.models.institution import Institution
from app.models.subject import Subject
from app.models.user import User, UserRole
from app.services import attendance_excel
from app.services.attendance import compute_student, percent


class AttendanceMathTests(unittest.TestCase):
    def test_percent_no_lectures_is_none(self):
        self.assertIsNone(percent(0, 0))

    def test_final_is_average_of_subject_percentages_not_pooled(self):
        # Subject 1: 1/1 = 100%. Subject 2: 1/3 = 33.33%. Mean = 66.67, pooled would be 50.
        sessions = {1: [10], 2: [20, 21, 22]}
        marks = {(10, 5): True, (20, 5): True, (21, 5): False, (22, 5): False}
        per_subject, final = compute_student([1, 2], sessions, marks, 5)
        self.assertEqual(per_subject[1]["percent"], 100.0)
        self.assertEqual(per_subject[2]["percent"], 33.33)
        self.assertEqual(final, 66.67)

    def test_subject_without_lectures_is_skipped_in_average(self):
        sessions = {1: [10], 2: []}
        per_subject, final = compute_student([1, 2], sessions, {(10, 5): True}, 5)
        self.assertIsNone(per_subject[2]["percent"])
        self.assertEqual(final, 100.0)

    def test_lecture_before_student_joined_is_not_counted(self):
        sessions = {1: [10, 11]}
        per_subject, final = compute_student([1], sessions, {(11, 5): True}, 5)
        self.assertEqual(per_subject[1]["held"], 1)
        self.assertEqual(final, 100.0)

    def test_no_lectures_at_all(self):
        _, final = compute_student([1], {}, {}, 5)
        self.assertIsNone(final)


class AttendanceApiTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.engine = create_engine(
            "sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool
        )
        Base.metadata.create_all(bind=cls.engine)
        cls.SessionLocal = sessionmaker(bind=cls.engine)

        def override_get_db():
            db = cls.SessionLocal()
            try:
                yield db
            finally:
                db.close()

        app.dependency_overrides[get_db] = override_get_db
        cls.client = TestClient(app)
        cls.tmp = tempfile.TemporaryDirectory()
        cls.patch = mock.patch.object(attendance_excel, "ATTENDANCE_DIR", Path(cls.tmp.name))
        cls.patch.start()

    @classmethod
    def tearDownClass(cls):
        cls.patch.stop()
        cls.tmp.cleanup()
        app.dependency_overrides.clear()
        Base.metadata.drop_all(bind=cls.engine)

    def _user(self, name, role, institution_id):
        u = User(
            full_name=name,
            email=f"{name.lower().replace(' ', '')}@example.com",
            hashed_password=get_password_hash("password123"),
            role=role,
            institution_id=institution_id,
        )
        self.db.add(u)
        self.db.commit()
        self.db.refresh(u)
        return u

    def _headers(self, user):
        return {"Authorization": f"Bearer {create_access_token(str(user.id), user.role.value)}"}

    def setUp(self):
        self.db = self.SessionLocal()
        for model in (
            AttendanceRecord,
            AttendanceSession,
            ClassroomStudent,
            ClassroomTeacher,
            Subject,
            Classroom,
            User,
            Institution,
        ):
            self.db.query(model).delete()
        self.db.commit()
        inst = Institution(name="Inst", code="INST")
        self.db.add(inst)
        self.db.commit()
        self.db.refresh(inst)

        self.class_teacher = self._user("Class Teacher", UserRole.CLASS_TEACHER, inst.id)
        self.subj_teacher = self._user("Subject Teacher", UserRole.SUBJECT_TEACHER, inst.id)
        self.other_teacher = self._user("Other Teacher", UserRole.SUBJECT_TEACHER, inst.id)
        self.alice = self._user("Alice", UserRole.STUDENT, inst.id)
        self.bob = self._user("Bob", UserRole.STUDENT, inst.id)

        self.classroom = Classroom(
            institution_id=inst.id,
            class_teacher_id=self.class_teacher.id,
            name="CS",
            code="CS1",
            join_code="ABCDE",
        )
        self.db.add(self.classroom)
        self.db.commit()
        self.db.refresh(self.classroom)
        for s in (self.alice, self.bob):
            self.db.add(
                ClassroomStudent(
                    classroom_id=self.classroom.id,
                    student_id=s.id,
                    status=MembershipStatus.APPROVED,
                )
            )
        self.maths = Subject(
            classroom_id=self.classroom.id,
            teacher_id=self.subj_teacher.id,
            name="Maths",
            code="MTH",
        )
        self.physics = Subject(
            classroom_id=self.classroom.id,
            teacher_id=self.class_teacher.id,
            name="Physics",
            code="PHY",
        )
        self.db.add_all([self.maths, self.physics])
        # Mirrors subjects.py _sync_classroom_teacher: assigned teachers get a row.
        for teacher in (self.subj_teacher, self.other_teacher):
            self.db.add(
                ClassroomTeacher(classroom_id=self.classroom.id, teacher_id=teacher.id)
            )
        self.db.commit()
        self.db.refresh(self.maths)
        self.db.refresh(self.physics)

    def tearDown(self):
        self.db.close()

    def _url(self, subject):
        return f"/api/v1/classrooms/{self.classroom.id}/subjects/{subject.id}/attendance"

    def _create(self, subject, teacher, day="2026-03-02", alice=True, bob=False):
        return self.client.post(
            self._url(subject),
            headers=self._headers(teacher),
            json={
                "held_on": day,
                "title": "",
                "marks": [
                    {"student_id": self.alice.id, "present": alice},
                    {"student_id": self.bob.id, "present": bob},
                ],
            },
        )

    def test_subject_teacher_can_mark_own_subject(self):
        res = self._create(self.maths, self.subj_teacher)
        self.assertEqual(res.status_code, 201, res.text)

    def test_teacher_cannot_mark_someone_elses_subject(self):
        # A subject teacher cannot touch another teacher's (or the class teacher's) subject.
        self.assertEqual(self._create(self.physics, self.subj_teacher).status_code, 403)
        self.assertEqual(self._create(self.maths, self.other_teacher).status_code, 403)

    def test_class_teacher_can_create_edit_delete_subject_teachers_attendance(self):
        res = self._create(self.maths, self.class_teacher, alice=True, bob=False)
        self.assertEqual(res.status_code, 201, res.text)
        by_subject_teacher = self._create(self.maths, self.subj_teacher, day="2026-03-03").json()
        h = self._headers(self.class_teacher)
        put = self.client.put(
            f"/api/v1/attendance/sessions/{by_subject_teacher['id']}",
            headers=h,
            json={"marks": [{"student_id": self.bob.id, "present": True}]},
        )
        self.assertEqual(put.status_code, 200, put.text)
        self.assertTrue(
            self.client.get(self._url(self.maths), headers=h).json()["can_edit"]
        )
        self.assertEqual(
            self.client.delete(f"/api/v1/attendance/sessions/{by_subject_teacher['id']}", headers=h).status_code,
            204,
        )

    def test_student_cannot_write(self):
        self.assertEqual(self._create(self.maths, self.alice).status_code, 403)

    def test_duplicate_lecture_conflicts(self):
        self.assertEqual(self._create(self.maths, self.subj_teacher).status_code, 201)
        self.assertEqual(self._create(self.maths, self.subj_teacher).status_code, 409)

    def test_non_member_student_rejected(self):
        res = self.client.post(
            self._url(self.maths),
            headers=self._headers(self.subj_teacher),
            json={"held_on": "2026-03-02", "marks": [{"student_id": 9999, "present": True}]},
        )
        self.assertEqual(res.status_code, 400)

    def test_class_teacher_reads_all_subject_teacher_reads_own(self):
        self._create(self.maths, self.subj_teacher)
        self.assertEqual(
            self.client.get(self._url(self.maths), headers=self._headers(self.class_teacher)).status_code,
            200,
        )
        self.assertEqual(
            self.client.get(self._url(self.maths), headers=self._headers(self.other_teacher)).status_code,
            403,
        )
        read = self.client.get(self._url(self.maths), headers=self._headers(self.class_teacher)).json()
        self.assertTrue(read["can_edit"])
        own = self.client.get(self._url(self.maths), headers=self._headers(self.subj_teacher)).json()
        self.assertTrue(own["can_edit"])

    def test_students_cannot_use_teacher_reads(self):
        self.assertEqual(
            self.client.get(self._url(self.maths), headers=self._headers(self.alice)).status_code, 403
        )
        base = f"/api/v1/classrooms/{self.classroom.id}/attendance"
        self.assertEqual(
            self.client.get(f"{base}/summary", headers=self._headers(self.alice)).status_code, 403
        )
        self.assertEqual(
            self.client.get(f"{base}/export.xlsx", headers=self._headers(self.alice)).status_code, 403
        )

    def test_student_sees_only_own_data_and_final_average(self):
        self._create(self.maths, self.subj_teacher, alice=True, bob=False)
        self._create(self.physics, self.class_teacher, alice=False, bob=True)
        res = self.client.get(
            f"/api/v1/me/classrooms/{self.classroom.id}/attendance",
            headers=self._headers(self.alice),
        )
        self.assertEqual(res.status_code, 200)
        body = res.json()
        by_code = {s["subject_code"]: s for s in body["subjects"]}
        self.assertEqual(by_code["MTH"]["percent"], 100.0)
        self.assertEqual(by_code["PHY"]["percent"], 0.0)
        self.assertEqual(body["final_percent"], 50.0)
        self.assertNotIn("Bob", res.text)

    def test_teacher_cannot_use_student_endpoint(self):
        res = self.client.get(
            f"/api/v1/me/classrooms/{self.classroom.id}/attendance",
            headers=self._headers(self.class_teacher),
        )
        self.assertEqual(res.status_code, 403)

    def test_workbook_follows_create_edit_delete(self):
        created = self._create(self.maths, self.subj_teacher, alice=True, bob=False).json()
        path = Path(self.tmp.name) / f"classroom_{self.classroom.id}.xlsx"
        self.assertTrue(path.exists())
        wb = load_workbook(path)
        self.assertIn("Summary", wb.sheetnames)
        self.assertIn("MTH", wb.sheetnames)
        rows = {r[0]: r for r in wb["MTH"].iter_rows(min_row=2, values_only=True)}
        self.assertEqual(rows["Alice"][2], "P")
        self.assertEqual(rows["Bob"][2], "A")

        res = self.client.put(
            f"/api/v1/attendance/sessions/{created['id']}",
            headers=self._headers(self.subj_teacher),
            json={
                "marks": [
                    {"student_id": self.alice.id, "present": False},
                    {"student_id": self.bob.id, "present": True},
                ]
            },
        )
        self.assertEqual(res.status_code, 200, res.text)
        rows = {r[0]: r for r in load_workbook(path)["MTH"].iter_rows(min_row=2, values_only=True)}
        self.assertEqual(rows["Alice"][2], "A")
        self.assertEqual(rows["Bob"][2], "P")

        res = self.client.delete(
            f"/api/v1/attendance/sessions/{created['id']}", headers=self._headers(self.subj_teacher)
        )
        self.assertEqual(res.status_code, 204)
        rows = {r[0]: r for r in load_workbook(path)["MTH"].iter_rows(min_row=2, values_only=True)}
        self.assertEqual(rows["Alice"][-2:], (0, None))  # Held = 0, % blank; no lecture columns left

    def test_export_class_teacher_full_subject_teacher_scoped(self):
        self._create(self.maths, self.subj_teacher)
        self._create(self.physics, self.class_teacher)
        url = f"/api/v1/classrooms/{self.classroom.id}/attendance/export.xlsx"
        full = self.client.get(url, headers=self._headers(self.class_teacher))
        self.assertEqual(full.status_code, 200)
        scoped = self.client.get(url, headers=self._headers(self.subj_teacher))
        self.assertEqual(scoped.status_code, 200)
        import io

        self.assertEqual(
            set(load_workbook(io.BytesIO(full.content)).sheetnames), {"Summary", "MTH", "PHY"}
        )
        self.assertEqual(
            set(load_workbook(io.BytesIO(scoped.content)).sheetnames), {"Summary", "MTH"}
        )


if __name__ == "__main__":
    unittest.main()
